"""Gera os templates .docx (com tags Jinja/docxtpl) em templates/docs/. Rode uma vez: `python -m licitabot.docs.build_templates`.

Os arquivos gerados podem ser abertos no Word e ajustados (logo, fontes, cabeçalho) desde que as tags sejam mantidas.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from licitabot.config import get_settings


def _base() -> Document:
    d = Document()
    for s in d.sections:
        s.top_margin = Cm(2.0)
        s.bottom_margin = Cm(2.0)
        s.left_margin = Cm(2.5)
        s.right_margin = Cm(2.0)
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    return d


def _cabecalho(d: Document) -> None:
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("{{ empresa.razao_social }}")
    r.bold = True
    r.font.size = Pt(14)
    p2 = d.add_paragraph("CNPJ {{ empresa.cnpj }} · {{ empresa.endereco }} · {{ empresa.cidade }}/{{ empresa.uf }} · CEP {{ empresa.cep }}")
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.runs[0].font.size = Pt(9)
    p3 = d.add_paragraph("{{ empresa.telefone }} · {{ empresa.email }} · {{ empresa.site }}")
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.runs[0].font.size = Pt(9)
    d.add_paragraph()


def _destinatario(d: Document) -> None:
    d.add_paragraph("À {{ op.orgao_nome }}")
    d.add_paragraph("Ref.: {{ numero_edital }}")
    d.add_paragraph("Objeto: {{ objeto }}")
    d.add_paragraph()


def _assinatura(d: Document) -> None:
    d.add_paragraph()
    d.add_paragraph("{{ empresa.cidade }}, {{ data_extenso }}.")
    d.add_paragraph()
    d.add_paragraph()
    p = d.add_paragraph("_______________________________________________")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for linha in ("{{ empresa.representante.nome }}", "{{ empresa.representante.cargo }} · CPF {{ empresa.representante.cpf }}", "{{ empresa.razao_social }} · CNPJ {{ empresa.cnpj }}"):
        q = d.add_paragraph(linha)
        q.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _tabela_loop(d: Document, cabecalho: list[str], for_expr: str, celulas: list[str]):
    """Tabela com cabeçalho + linha `{%tr for %}` + linha de conteúdo + linha `{%tr endfor %}` (padrão docxtpl)."""
    tabela = d.add_table(rows=4, cols=len(cabecalho))
    tabela.style = "Table Grid"
    for i, c in enumerate(cabecalho):
        cell = tabela.rows[0].cells[i]
        cell.text = c
        cell.paragraphs[0].runs[0].bold = True
    tabela.rows[1].cells[0].text = "{%tr " + for_expr + " %}"
    for i, c in enumerate(celulas):
        tabela.rows[2].cells[i].text = c
    tabela.rows[3].cells[0].text = "{%tr endfor %}"
    return tabela


def build_declaracao(dest: Path) -> None:
    d = _base()
    _cabecalho(d)
    t = d.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("{{ titulo }}")
    r.bold = True
    r.font.size = Pt(12)
    d.add_paragraph()
    d.add_paragraph("Ref.: {{ numero_edital }} — {{ op.orgao_nome }}")
    d.add_paragraph()
    p = d.add_paragraph("{{ texto }}")
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _assinatura(d)
    d.save(dest)


def build_proposta_comercial(dest: Path) -> None:
    d = _base()
    _cabecalho(d)
    t = d.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("PROPOSTA COMERCIAL")
    r.bold = True
    r.font.size = Pt(13)
    _destinatario(d)
    d.add_paragraph(
        "Apresentamos nossa proposta comercial para o objeto acima, em conformidade com o edital e seus anexos, conforme discriminado a seguir:"
    ).paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    tabela = _tabela_loop(
        d,
        ["Item", "Descrição", "Unid.", "Qtd.", "Marca/Modelo", "Valor unitário (R$)", "Valor total (R$)"],
        "for it in itens",
        ["{{ it.numero_item }}", "{{ it.descricao }}", "{{ it.unidade }}", "{{ it.quantidade_fmt }}", "{{ it.marca_modelo }}", "{{ it.valor_unit_fmt }}", "{{ it.valor_total_fmt }}"],
    )
    tabela.alignment = WD_TABLE_ALIGNMENT.CENTER
    d.add_paragraph()
    p = d.add_paragraph()
    p.add_run("VALOR TOTAL DA PROPOSTA: R$ {{ total_fmt }} ({{ total_extenso }})").bold = True
    d.add_paragraph()
    for linha_txt in (
        "Validade da proposta: {{ validade_dias }} ({{ validade_dias_num }}) dias, contados da data de abertura da sessão pública.",
        "Prazo de execução/entrega: {{ prazo_execucao }}.",
        "Nos preços propostos estão incluídos todos os custos diretos e indiretos, tributos, encargos sociais, trabalhistas e previdenciários, "
        "fretes, seguros e quaisquer outras despesas necessárias à perfeita execução do objeto.",
        "Declaramos que a proposta foi elaborada de forma independente e que atende integralmente às especificações do Termo de Referência.",
        "Dados bancários: {{ empresa.banco.banco }} · Agência {{ empresa.banco.agencia }} · Conta {{ empresa.banco.conta }} ({{ empresa.banco.tipo }}).",
        "Representante legal para assinatura do contrato: {{ empresa.representante.nome }}, {{ empresa.representante.cargo }}, CPF {{ empresa.representante.cpf }}, "
        "e-mail {{ empresa.representante.email }}, telefone {{ empresa.representante.telefone }}.",
    ):
        q = d.add_paragraph(linha_txt, style="List Bullet")
        q.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    d.add_paragraph()
    d.add_paragraph("{% if observacoes %}Observações: {{ observacoes }}{% endif %}")
    _assinatura(d)
    d.save(dest)


def build_proposta_tecnica(dest: Path) -> None:
    d = _base()
    _cabecalho(d)
    t = d.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("PROPOSTA TÉCNICA")
    r.bold = True
    r.font.size = Pt(13)
    _destinatario(d)
    d.add_heading("Resumo executivo", level=1)
    d.add_paragraph("{{ resumo_executivo }}").paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    d.add_paragraph("{%p for s in secoes %}")
    d.add_heading("{{ s.titulo }}", level=1)
    d.add_paragraph("{{ s.texto }}").paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    d.add_paragraph("{%p endfor %}")
    d.add_paragraph("{%p if premissas %}")
    d.add_heading("Premissas", level=1)
    d.add_paragraph("{%p for pr in premissas %}")
    d.add_paragraph("{{ pr }}", style="List Bullet")
    d.add_paragraph("{%p endfor %}")
    d.add_paragraph("{%p endif %}")
    _assinatura(d)
    d.save(dest)


def build_planilha_custos(dest: Path) -> None:
    d = _base()
    _cabecalho(d)
    t = d.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("PLANILHA DE COMPOSIÇÃO DE CUSTOS E FORMAÇÃO DE PREÇOS")
    r.bold = True
    _destinatario(d)
    _tabela_loop(
        d,
        ["Item", "Composição", "Qtd. / Horas", "Custo unitário (R$)", "Total (R$)"],
        "for it in itens",
        ["{{ it.numero_item }}", "{{ it.justificativa_preco }}", "{{ it.quantidade_fmt }}", "{{ it.valor_unit_fmt }}", "{{ it.valor_total_fmt }}"],
    )
    d.add_paragraph()
    d.add_paragraph("Parâmetros: margem {{ margem_pct }}% · tributos embutidos {{ impostos_pct }}% · valor/hora de referência R$ {{ valor_hora_fmt }}.")
    d.add_paragraph("Total: R$ {{ total_fmt }}")
    _assinatura(d)
    d.save(dest)


def build_indice(dest: Path) -> None:
    d = _base()
    _cabecalho(d)
    t = d.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("RELAÇÃO DE DOCUMENTOS APRESENTADOS")
    r.bold = True
    _destinatario(d)
    _tabela_loop(d, ["Nº", "Documento", "Fase", "Arquivo"], "for dcm in documentos", ["{{ dcm.ordem }}", "{{ dcm.titulo }}", "{{ dcm.fase }}", "{{ dcm.arquivo }}"])
    _assinatura(d)
    d.save(dest)


def build_all(dest_dir: Path | None = None) -> list[Path]:
    dest_dir = dest_dir or (get_settings().templates_path / "docs")
    dest_dir.mkdir(parents=True, exist_ok=True)
    out = []
    for nome, fn in (
        ("declaracao.docx", build_declaracao),
        ("proposta_comercial.docx", build_proposta_comercial),
        ("proposta_tecnica.docx", build_proposta_tecnica),
        ("planilha_custos.docx", build_planilha_custos),
        ("indice_documentos.docx", build_indice),
    ):
        p = dest_dir / nome
        if not p.exists():  # não sobrescreve ajustes manuais do usuário
            fn(p)
        out.append(p)
    return out


if __name__ == "__main__":
    for p in build_all():
        print(p)
